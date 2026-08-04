#!/usr/bin/env python3
"""
Refresh the bundled workbook from a downloaded copy.

The project ships the workbook rather than fetching it live, because the
SharePoint link requires a Northeastern sign-in (see docs/onedrive.md). This
script makes the manual refresh safe.

Its real job is guarding against picking the wrong file. There are several
similarly-named workbooks in ~/Downloads representing different passes
(104 / 134 / 194 facilities), so this counts the rows in each candidate,
shows you what changed, and refuses to silently install a smaller dataset.

    python3 scripts/refresh.py              # inspect, then prompt
    python3 scripts/refresh.py --list       # just show candidates
    python3 scripts/refresh.py --yes        # no prompt (must not shrink)
    python3 scripts/refresh.py --file ~/Downloads/other.xlsx
"""

from __future__ import annotations

import argparse
import glob
import os
import shutil
import subprocess
import sys
from datetime import datetime

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required:  python3 -m pip install openpyxl")

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
BUNDLED = os.path.join(ROOT, "data", "master_dataset_1_values_2.xlsx")
DEFAULT_SRC = os.path.expanduser("~/Downloads")

SHEET = "Facilities"
CONSOLIDATED = "Consolidated (all layers)"


def inspect(path):
    """Facility/operator/country counts without a full parse."""
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as exc:
        return {"path": path, "error": str(exc)}
    info = {"path": path, "sheets": wb.sheetnames,
            "mtime": datetime.fromtimestamp(os.path.getmtime(path))}
    if SHEET not in wb.sheetnames:
        info["error"] = f'no "{SHEET}" sheet'
        wb.close()
        return info

    ws = wb[SHEET]
    headers = [str(c.value).strip() if c.value else "" for c in next(ws.iter_rows(max_row=1))]
    try:
        i_id = headers.index("Facility ID")
        i_country = headers.index("Country")
        i_op = headers.index("Operator")
    except ValueError as exc:
        info["error"] = f"missing column: {exc}"
        wb.close()
        return info

    ids, countries, operators = set(), set(), set()
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or i_id >= len(row) or not row[i_id]:
            continue
        ids.add(str(row[i_id]).strip())
        if i_country < len(row) and row[i_country]:
            countries.add(str(row[i_country]).strip())
        if i_op < len(row) and row[i_op]:
            operators.add(str(row[i_op]).strip())

    info.update(facilities=len(ids), countries=len(countries), operators=len(operators),
                has_consolidated=CONSOLIDATED in wb.sheetnames)
    wb.close()
    return info


def describe(info, label=""):
    if "error" in info:
        return f"  {label}{os.path.basename(info['path'])}: ERROR — {info['error']}"
    flag = "" if info["has_consolidated"] else "  [!] no consolidated sheet"
    return (f"  {label}{os.path.basename(info['path']):38} "
            f"{info['facilities']:4} facilities · {info['countries']:2} countries · "
            f"{info['operators']:3} operators · {info['mtime']:%Y-%m-%d %H:%M}{flag}")


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--src", default=DEFAULT_SRC, help="folder to scan (default ~/Downloads)")
    ap.add_argument("--file", help="use this workbook explicitly, skipping the scan")
    ap.add_argument("--list", action="store_true", help="show candidates and exit")
    ap.add_argument("--yes", action="store_true", help="don't prompt (refuses to shrink the dataset)")
    args = ap.parse_args()

    current = inspect(BUNDLED) if os.path.exists(BUNDLED) else None
    print("currently bundled:")
    print(describe(current) if current else "  (none)")
    print()

    if args.file:
        candidates = [inspect(os.path.expanduser(args.file))]
    else:
        found = sorted(glob.glob(os.path.join(args.src, "master_dataset*.xlsx")))
        # Skip Excel's lock files for open workbooks.
        found = [f for f in found if not os.path.basename(f).startswith("~$")]
        if not found:
            sys.exit(f"no master_dataset*.xlsx found in {args.src}")
        candidates = [inspect(f) for f in found]

    usable = [c for c in candidates if "error" not in c]
    print(f"candidates in {args.src if not args.file else 'given path'}:")
    for c in candidates:
        best = usable and c is max(usable, key=lambda x: x["facilities"])
        print(describe(c, "→ " if best else "  "))
    print()

    if args.list:
        return
    if not usable:
        sys.exit("no usable candidate")

    pick = max(usable, key=lambda x: x["facilities"])
    if not pick["has_consolidated"]:
        sys.exit(f'refusing: {os.path.basename(pick["path"])} has no "{CONSOLIDATED}" sheet')

    if os.path.abspath(pick["path"]) == os.path.abspath(BUNDLED):
        print("the bundled copy is already the largest candidate — nothing to do")
        return

    if current and "error" not in current:
        delta = pick["facilities"] - current["facilities"]
        if delta < 0:
            msg = (f"that file has {abs(delta)} FEWER facilities "
                   f"({pick['facilities']} vs {current['facilities']})")
            if args.yes:
                sys.exit(f"refusing: {msg}. Re-run without --yes to confirm interactively.")
            print(f"warning: {msg}")
            if input("proceed anyway? [y/N] ").strip().lower() != "y":
                return
        elif delta == 0:
            print(f"same facility count ({pick['facilities']}) — refreshing anyway "
                  f"(cell values may still have changed)")
        else:
            print(f"+{delta} facilities ({current['facilities']} → {pick['facilities']})")

    if not args.yes:
        if input(f"install {os.path.basename(pick['path'])}? [y/N] ").strip().lower() != "y":
            return

    shutil.copy2(pick["path"], BUNDLED)
    print(f"\ncopied → {BUNDLED}")

    # Regenerate the flat exports and show the data-quality report, so a refresh
    # never silently changes capacity coverage without you seeing it.
    # Flush first: when stdout is a pipe our prints are buffered while the child
    # writes straight to the fd, which otherwise interleaves the output backwards.
    sys.stdout.flush()
    for cmd in (["--check"], []):
        subprocess.run([sys.executable, os.path.join(HERE, "export_data.py")] + cmd, check=True)

    print("\nReload the app to see it. To publish: git add -A && git commit && git push")


if __name__ == "__main__":
    main()
